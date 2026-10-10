"""Common-format content/location/reparse qualification via actual candidate Core.

Project-authored synthetic DOCX/XLSX/PPTX/text-PDF materials. No direct SQLite
writes, installed UI, source-application rendering, OCR or model inference claim.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import subprocess
import time
import uuid
from pathlib import Path

import aaos01_office_runtime_loop as office

REPO = Path(__file__).resolve().parents[2]
FORMATS = ("docx", "xlsx", "pptx", "pdf")
EXPECTED = {
    "docx": ("DOCX中文标题", "正文甲", "表格乙", "重复结构中文"),
    "xlsx": ("表格中文证据", "=A2*2", "第二页中文", "合并中文"),
    "pptx": ("幻灯中文一", "幻灯中文二", "备注中文", "类别中文"),
    "pdf": ("PDF中文页一", "PDF中文页二"),
}

GENERATOR = r"""
import json,sys,zipfile,base64
from pathlib import Path
import fitz,openpyxl,pptx
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
root=Path(sys.argv[1]);root.mkdir(exist_ok=False)
xml='<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>DOCX中文标题</w:t></w:r></w:p><w:p><w:r><w:t>正文甲</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>表格乙</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>42</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p><w:r><w:t>重复结构中文</w:t></w:r></w:p><w:p><w:r><w:t>重复结构中文</w:t></w:r></w:p></w:body></w:document>'
xml=xml.replace('<w:document xmlns:w=', '<w:document xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:w=')
xml=xml.replace('</w:body>', '<w:sectPr><w:headerReference w:type="default" r:id="rHeader"/></w:sectPr></w:body>')
with zipfile.ZipFile(root/'complex.docx','w') as z:
 z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Default Extension="png" ContentType="image/png"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
 z.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rMain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
 z.writestr('word/_rels/document.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rHeader" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/><Relationship Id="rStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship Id="rImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/></Relationships>')
 z.writestr('word/document.xml',xml)
 z.writestr('word/header1.xml','<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>HEADER_OMITTED_中文</w:t></w:r></w:p></w:hdr>')
 z.writestr('word/styles.xml','<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>')
 z.writestr('word/media/image1.png',base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jq1kAAAAASUVORK5CYII='))
b=openpyxl.Workbook();s=b.active;s.title='证据';s.append(['表格中文证据','数值']);s.append([12,'=A2*2']);s.merge_cells('A4:B4');s['A4']='合并中文';b.create_sheet('第二页')['A1']='第二页中文';b.save(root/'complex.xlsx')
p=pptx.Presentation()
for text in ('幻灯中文一','幻灯中文二'):
 s=p.slides.add_slide(p.slide_layouts[6]);s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text=text
 s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text='重复结构中文'
 s.notes_slide.notes_text_frame.text='备注中文'
c=CategoryChartData();c.categories=['类别中文','类别乙'];c.add_series('缓存系列',[12,24]);p.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(5), Inches(3),c);p.save(root/'complex.pptx')
d=fitz.open()
for text in ('PDF中文页一','PDF中文页二'):
 page=d.new_page();page.insert_text((72,90),text,fontname='china-s');page.insert_text((72,130),'内容中文42',fontname='china-s')
d.save(root/'complex.pdf');d.close()
print(json.dumps({'formats':['docx','xlsx','pptx','pdf'],'materials':'SYNTHETIC_AUTHORED','engines':{'PyMuPDF':fitz.VersionBind,'openpyxl':openpyxl.__version__,'python-pptx':pptx.__version__}},ensure_ascii=False))
"""


def validate_outputs(extension, snapshot):
    assert snapshot["job"]["body"]["state"] == "succeeded", snapshot["job"]
    for key in ("text", "document_structure", "loss_report", "quality"):
        assert snapshot[key]["status"] == 200, snapshot[key]
    for key in ("text", "document_structure", "loss_report"):
        output = snapshot[key]["body"]
        raw = output["content"].encode("utf-8")
        assert output["metadata"]["sha256"] == hashlib.sha256(raw).hexdigest()
        assert output["metadata"]["byte_length"] == len(raw)
    text = snapshot["text"]["body"]["content"]
    compact = "".join(text.split())
    assert all("".join(marker.split()) in compact for marker in EXPECTED[extension]), text
    loss = json.loads(snapshot["loss_report"]["body"]["content"])
    assert loss["engine"] == ("pymupdf-native-pdf" if extension == "pdf" else "python-worker-office")
    assert isinstance(loss.get("losses"), list) and isinstance(loss.get("loss_note"), str)
    params = loss["params"]
    if extension == "docx":
        assert params["headers"] and params["media_files"] == 1
        assert "HEADER_OMITTED" not in text
        assert "header" in loss["loss_note"].lower() and "inventoried" in loss["loss_note"]
    elif extension == "xlsx":
        assert params["sheets"] == 2 and params["formula_cells"] == 1
        assert "not" in loss["loss_note"].lower() and "calculation" in loss["loss_note"].lower()
    elif extension == "pptx":
        assert params["slides"] == 2 and params["charts_with_cached_values"] == 1
        assert "cached" in loss["loss_note"] and "never recomputed" in loss["loss_note"]
    else:
        assert params["pages"] == 2 and params["pages_without_text"] == []
    positions = []
    reported = json.loads(snapshot["document_structure"]["body"]["content"]) if extension == "pdf" else params["worker_structure"]
    for position in reported:
        start, end = position["char_start"], position["char_end"]
        assert 0 <= start < end <= len(text), position
        if text[start:end].strip() and position.get("path"):
            positions.append(position)
    assert len(positions) >= 2, "composite structure not reported"
    return text, loss, positions


def create_location(call, source, job, snapshot, position, native_pdf=False):
    text = snapshot["text"]["body"]["content"]
    excerpt = text[position["char_start"]:position["char_end"]]
    locator = {"type": "pdf_line" if native_pdf else "worker_structure", "job_id": job,
               "attempt": snapshot["job"]["body"]["attempt"],
               "kind": position["kind"], "path": position["path"]}
    if native_pdf:
        locator.update({"char_start": position["char_start"], "char_end": position["char_end"],
                        "result_sha256": snapshot["document_structure"]["body"]["metadata"]["sha256"]})
    body = {"revision": source["sha256"], "checksum": hashlib.sha256(excerpt.encode()).hexdigest(),
            "position": json.dumps(locator)}
    path = f"/api/v1/sources/{source['source_id']}/anchors"
    anchor = call("POST", path, body, 201)
    assert anchor["location_status"] == "located"
    call("POST", path, {**body, "checksum": "0" * 64}, 400)
    result = call("GET", path + f"/{anchor['anchor_id']}/resolve")
    assert result["status"] == "CURRENT" and result["scope"] == "locator_provenance_only", result
    return {"anchor": anchor, "resolution": result}


def run_job(call, client, base, token, source, extension):
    job = "common-" + uuid.uuid4().hex
    call("POST", "/api/v1/jobs", {"job_id": job, "kind": "pdf" if extension == "pdf" else "office",
                                  "input_ref": source["source_id"]}, 202)
    call("POST", f"/api/v1/jobs/{job}/executions", {"deadline_ms": 60000}, 202,
         headers={"idempotency-key": job})
    deadline = time.monotonic() + 65
    while True:
        state = call("GET", f"/api/v1/jobs/{job}")
        if state["state"] in ("succeeded", "failed", "cancelled", "rejected"):
            break
        assert time.monotonic() < deadline, "worker deadline expired"
        time.sleep(.1)
    snapshot = office.capture(client, base, token, job)
    _, _, positions = validate_outputs(extension, snapshot)
    second = next((p for p in positions if p["kind"] == "table_row"), None) if extension == "docx" else next(
        (p for p in positions if p["path"][0] != positions[0]["path"][0]), None)
    assert second is not None, "second structural unit not available for location"
    return {"job_id": job, "snapshot": snapshot, "location_qualification": "CURRENT_REVALIDATED",
            "locations": [create_location(call, source, job, snapshot, p, native_pdf=extension == "pdf")
                          for p in (positions[0], second)]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--candidate", type=Path, required=True)
    args = parser.parse_args()
    candidate = args.candidate.resolve()
    dev = office.load("common_dev", REPO / "scripts/runtime/dev.py")
    paths = dev.layout(REPO)
    dev.prepare(paths)
    work = paths["run"]
    launcher = office.load("common_launcher", REPO / "scripts/release/backend_launcher.py")
    client = office.load("common_client", REPO / "shared/core_client.py")
    receipt = {"ok": False, "materials": "SYNTHETIC_AUTHORED", "execution": "ACTUAL_CANDIDATE_CORE_WORKERS",
               "limits": ["No installed UI/source-application navigation", "Text PDF only; OCR separate",
                          "Formula text/cached chart data are not live calculation", "No source package reconstruction",
                          "Independent reparse uses a new job; old pinned job remains valid, not superseded"],
               "qualification": "NOT_EXECUTED",
               "formats": []}
    child = None
    try:
        profile = launcher.load_profile(candidate)
        dirty, patch_sha = dev.worktree_identity(REPO)
        receipt["source"] = {"commit": dev.git(REPO, "rev-parse", "HEAD"), "dirty": dirty,
                             "patch_sha256": patch_sha, "changes": office.source_changes(dev)}
        receipt["core"] = office.identity(candidate / "core/archeaxis-api.exe")
        receipt["manifest"] = office.identity(candidate / "backend-runtime-manifest.json")
        receipt["interpreter"] = office.identity(profile["python"])
        receipt["worker_profile"] = office.identity(candidate / "worker-profile.json")
        material = work / "materials"
        generated = subprocess.run([str(profile["python"]), "-B", "-I", "-c", GENERATOR, str(material)],
                                   capture_output=True, text=True, encoding="utf-8", timeout=60)
        receipt["generator"] = {"exit": generated.returncode, "stdout": generated.stdout, "stderr": generated.stderr}
        assert generated.returncode == 0, "candidate fixture generation failed"
        child, base, token, _ = office.start(candidate, work, launcher)

        def call(method, path, body=None, expected=200, headers=None):
            status, result = client.call(base, method, path, token, body, extra_headers=headers)
            assert status == expected, {"path": path, "status": status, "result": result}
            return result

        for extension in FORMATS:
            raw = (material / f"complex.{extension}").read_bytes()
            source = call("POST", "/api/v1/imports", client.import_request(f"complex.{extension}", raw), 202)
            assert source["sha256"] == hashlib.sha256(raw).hexdigest()
            original_path = f"/api/v1/sources/{source['source_id']}/original"
            original = call("GET", original_path)
            assert base64.b64decode(original["content_base64"]) == raw
            first = run_job(call, client, base, token, source, extension)
            anchor_path = f"/api/v1/sources/{source['source_id']}/anchors"
            historical = call("GET", anchor_path)
            second = run_job(call, client, base, token, source, extension)
            assert first["job_id"] != second["job_id"]
            for key in ("text", "document_structure"):
                assert first["snapshot"][key] == second["snapshot"][key], "independent reparse changed projection"
            after = call("GET", anchor_path)
            assert all(row in after["anchors"] for row in historical["anchors"]), "reparse rewrote historical anchors"
            assert call("GET", original_path) == original
            record = {"extension": extension, "input": office.identity(material / f"complex.{extension}"),
                      "source": source, "original_sha": hashlib.sha256(raw).hexdigest(),
                      "first": first, "reparse": second, "anchors": after}
            receipt["formats"].append(record)
        launcher.stop(child)
        child = None
        child, base, token, _ = office.start(candidate, work, launcher)
        for record in receipt["formats"]:
            source = record["source"]
            raw = (material / f"complex.{record['extension']}").read_bytes()
            assert base64.b64decode(call("GET", f"/api/v1/sources/{source['source_id']}/original")["content_base64"]) == raw
            anchor_path = f"/api/v1/sources/{source['source_id']}/anchors"
            assert call("GET", anchor_path) == record["anchors"]
            for phase in ("first", "reparse"):
                run = record[phase]
                assert office.capture(client, base, token, run["job_id"]) == run["snapshot"]
                for located in run["locations"]:
                    actual = call("GET", anchor_path + f"/{located['anchor']['anchor_id']}/resolve")
                    assert actual == located["resolution"], "restart changed independent resolution"
            record["restart_equal"] = True
        receipt["ok"] = True
        receipt["qualification"] = "INTEGRATED_SELECTED_CONTENT_LOCATION_REPARSE_RESTART"
    except Exception as error:
        receipt["error"] = f"{type(error).__name__}: {error}"
    finally:
        if child is not None:
            try:
                launcher.stop(child)
            except Exception as error:
                receipt["ok"] = False
                receipt["cleanup_error"] = f"{type(error).__name__}: {error}"
        output = paths["artifacts"] / "common-formats.json"
        output.write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"ok": receipt["ok"], "receipt": str(output), "error": receipt.get("error")}))
    return 0 if receipt["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
