"""Real synthetic OOXML duplicate-text/multiple-chart locator regression; no Core PASS claim."""
import importlib.util,os,tempfile,unittest,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
BASE=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('candidate_office',BASE/'services/python-workers/document/worker_office.py')
worker=importlib.util.module_from_spec(spec);spec.loader.exec_module(worker)
class PptxLocationIdentity(unittest.TestCase):
 def test_two_identical_unicode_shapes_have_distinct_package_identities_and_spans(self):
  with tempfile.TemporaryDirectory(dir=os.environ['ARCHEAXIS_RUN_ROOT']) as temp:
   p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6]);text='同一段 Unicode 中文'
   shapes=[slide.shapes.add_textbox(0,i*400000,2000000,300000) for i in range(2)]
   for shape in shapes:shape.text_frame.text=text
   path=Path(temp)/'duplicate.pptx';p.save(path);out=worker.extract(str(path));entries=out['structure']
   self.assertEqual(len(entries),2);self.assertEqual(len({tuple(e['path']) for e in entries}),2)
   self.assertEqual([e['shape_id'] for e in entries],[s.shape_id for s in shapes])
   for e,s in zip(entries,shapes):
    self.assertEqual(e['path'],['slide-1',f'shape-{s.shape_id}','slide']);self.assertEqual(out['text'][e['char_start']:e['char_end']],text)
   self.assertGreater(entries[1]['char_start'],entries[0]['char_end'])
   self.assertEqual(worker.extract(str(path))['structure'],entries)
 def test_two_charts_use_real_distinct_relationships_and_chart_parts(self):
  with tempfile.TemporaryDirectory(dir=os.environ['ARCHEAXIS_RUN_ROOT']) as temp:
   p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6])
   for i in range(2):
    data=CategoryChartData();data.categories=['Q1','Q2'];data.add_series('Revenue',(41.,92.5));slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,i*3000000,0,2000000,2000000,data)
   path=Path(temp)/'charts.pptx';p.save(path);out=worker.extract(str(path));entries=[e for e in out['structure'] if e['kind']=='slide_chart']
   self.assertEqual(len(entries),2);self.assertEqual(len({tuple(e['path']) for e in entries}),2)
   with zipfile.ZipFile(path) as z:
    rels=ET.fromstring(z.read('ppt/slides/_rels/slide1.xml.rels'));actual={r.get('Id'):r.get('Target') for r in rels if r.get('Type')==worker.CHART_REL_TYPE}
    for e in entries:
     self.assertIn(e['relationship_id'],actual);self.assertTrue(e['chart_part'].endswith(actual[e['relationship_id']].split('/')[-1]));self.assertIn(e['chart_part'],z.namelist())
     self.assertEqual(e['path'],['slide-1','relationship-'+e['relationship_id'],e['chart_part'],'slide_chart'])
     self.assertIn('Q1=41',out['text'][e['char_start']:e['char_end']])
   self.assertEqual(out['loss_receipt']['params']['charts_with_cached_values'],2)
 def test_project_owned_golden_text_shapes_are_individually_addressable(self):
  source=BASE/'tests/fixtures/golden/golden-pptx-anchor.pptx'
  p=Presentation(source);expected=[(i,s.shape_id,s.text_frame.text.strip()) for i,slide in enumerate(p.slides,1) for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
  out=worker.extract(str(source));entries=[e for e in out['structure'] if e['kind']=='slide']
  self.assertGreaterEqual(len(entries),2);self.assertEqual(len(entries),len(expected));self.assertEqual(len({tuple(e['path']) for e in entries}),len(entries))
  for e,(slide,shape,text) in zip(entries,expected):
   self.assertEqual(e['path'],[f'slide-{slide}',f'shape-{shape}','slide']);self.assertEqual(out['text'][e['char_start']:e['char_end']],text)
if __name__=='__main__':unittest.main()
