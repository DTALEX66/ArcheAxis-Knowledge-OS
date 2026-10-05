"""Real worker protocol regression for Windows long staging paths."""

import hashlib
import importlib.util
import io
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
TRANSPORT = ROOT / "services/python-workers/transport/text_ndjson.py"
SPEC = importlib.util.spec_from_file_location("windows_staging_transport", TRANSPORT)
transport = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(transport)


def fixture_io(path):
    """Create test assets independently of the production long-path helper."""
    return Path("\\\\?\\" + str(path.absolute())) if os.name == "nt" else path


@pytest.mark.skipif(os.name != "nt", reason="Windows MAX_PATH regression")
@pytest.mark.parametrize("kind", ["xlsx", "pptx", "text"])
def test_real_worker_extracts_and_replays_long_staging_paths(tmp_path, kind):
    if kind == "xlsx":
        from openpyxl import Workbook

        book = Workbook()
        book.active.title = "KnownSheet"
        book.active["A1"] = "Known long path cell"
        book.active["B1"] = "=1+2"
        book.active["A2"] = "\\\\?\\C:\\source\\original.txt"
        stream = io.BytesIO()
        book.save(stream)
        raw = stream.getvalue()
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        expected = "A1=Known long path cell"
    elif kind == "pptx":
        from pptx import Presentation

        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Known long path slide"
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "\\\\?\\C:\\source\\original.txt"
        stream = io.BytesIO()
        deck.save(stream)
        raw = stream.getvalue()
        media = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        expected = "Known long path slide"
    else:
        raw = b"Known long path text\n\\\\?\\C:\\source\\original.txt\n"
        media = "text/plain"
        expected = "Known long path text"
    staging = tmp_path / ("deep-staging-" * 8) / ("attempt-" * 8)
    assert len(str(staging)) > 260
    fixture_io(staging / "input").mkdir(parents=True)
    digest = hashlib.sha256(raw).hexdigest()
    fixture_io(staging / "input" / digest).write_bytes(raw)
    capability = "text.extract" if kind == "text" else "office.structure"
    request = {
        "schema": "archeaxis.worker-request/v1", "type": "job_request",
        "request_id": "long-path-request", "job_id": "long-path-job", "attempt": 1,
        "protocol_minor": 0, "capability": capability, "capability_version": "1",
        "deadline_ms": 30000, "parameters": {},
        "inputs": [{"uri": f"job://input/{digest}", "sha256": digest, "media_type": media}],
    }
    script = TRANSPORT if kind == "text" else ROOT / "services/python-workers/document/worker_office.py"
    responses = []
    for _ in range(2):
        result = subprocess.run(
            [sys.executable, "-B", str(script), "--staging-root", str(staging)],
            input=json.dumps(request) + "\n", capture_output=True, text=True,
            encoding="utf-8", timeout=30, check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        assert result.returncode == 0, result.stdout + result.stderr
        hello, response = [json.loads(line) for line in result.stdout.splitlines()]
        assert hello["capabilities"] == [capability]
        assert response["status"] == "succeeded" and response["error"] is None
        assert "\\\\?\\" not in result.stdout
        payloads = {}
        for output in response["outputs"]:
            assert output["uri"] == f"job://output/{output['sha256']}"
            content = fixture_io(staging / "output" / output["sha256"]).read_bytes()
            assert hashlib.sha256(content).hexdigest() == output["sha256"]
            assert len(content) == output["byte_length"]
            payloads[output["kind"]] = content
        assert expected in payloads["text"].decode("utf-8")
        assert "\\\\?\\C:\\source\\original.txt" in payloads["text"].decode("utf-8")
        anchors = json.loads(payloads["document_structure"])
        assert anchors and all(anchor["char_end"] > anchor["char_start"] for anchor in anchors)
        loss = json.loads(payloads["loss_report"])
        assert loss["engine"] and loss["engine_version"]
        if kind != "text":
            assert loss["losses"] and loss["params"]["worker_structure"]
        responses.append(response["outputs"])
    assert responses[0] == responses[1]


@pytest.mark.parametrize("value", ["E:/never-read", "E:never-read", "//server/share", "\\\\?\\C:\\never-read", "parent/../child", ".agents/private"])
def test_public_staging_paths_still_reject_protected_forms(value):
    with pytest.raises(transport.Rejected):
        transport.safe_path(Path(value), missing=True)


@pytest.mark.skipif(os.name != "nt", reason="Windows junction regression")
def test_long_staging_paths_still_reject_redirected_ancestors(tmp_path):
    donor = tmp_path / "donor"
    donor.mkdir()
    link = tmp_path / "redirect"
    subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(donor)],
                   check=True, capture_output=True)
    path = link / ("deep-staging-" * 8) / ("attempt-" * 8)
    assert len(str(path)) > 260
    with pytest.raises(transport.Rejected, match="reparse"):
        transport.safe_path(path, missing=True)


def test_route_contract_preserves_user_structure_and_loss_notes():
    prefix = "\\\\?\\C:\\source\\original.txt"
    structure = [{"kind": "paragraph", "path": [prefix], "char_start": 0, "char_end": len(prefix)}]
    result = transport._as_route_contract({
        "text": prefix, "structure": structure,
        "loss_receipt": {"loss_note": prefix, "params": {"source_note": prefix}},
    }, "office.structure")
    assert result["text"] == prefix
    assert result["loss_receipt"]["loss_note"] == prefix
    assert result["loss_receipt"]["params"]["source_note"] == prefix
    assert result["loss_receipt"]["params"]["worker_structure"] == structure
