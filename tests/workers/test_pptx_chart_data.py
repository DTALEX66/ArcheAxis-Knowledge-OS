"""F08: a PPTX chart must report what it holds, not merely that a chart exists.

The chart data is read straight from the package's chart part with the standard library, so no
new dependency enters the worker and no spreadsheet application is implied. Values are the
file's own cached results and the receipt says so; a series whose data lives only in a linked
workbook is named as carrying no cached values rather than guessed at.
"""

import importlib.util
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OFFICE = ROOT / "services/python-workers/document/worker_office.py"
CHART_PART = "ppt/charts/chart1.xml"


def _load_worker():
    spec = importlib.util.spec_from_file_location("worker_office", OFFICE)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _pptx_with_chart(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 400000, 200000)
    box.text_frame.text = "Quarterly revenue"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (41.0, 92.5))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 100000, 300000, 4000000, 2500000, data)
    presentation.save(str(path))
    return path


def _strip_caches(source: Path, target: Path) -> Path:
    """Rewrite the package with every chart cache emptied, as a chart linked to live data is."""
    with zipfile.ZipFile(source) as reader, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as writer:
        for item in reader.infolist():
            payload = reader.read(item.filename)
            if item.filename == CHART_PART:
                text = payload.decode("utf-8")
                for cache in ("strCache", "numCache"):
                    start = 0
                    while True:
                        open_tag = text.find(f"<c:{cache}", start)
                        if open_tag < 0:
                            break
                        close_tag = text.find(f"</c:{cache}>", open_tag)
                        head = text[open_tag:text.find(">", open_tag) + 1]
                        text = text[:open_tag] + head + text[close_tag:]
                        start = open_tag + len(head)
                payload = text.encode("utf-8")
            writer.writestr(item, payload)
    return target


class PptxChartDataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.worker = _load_worker()

    def _tmp(self):
        return tempfile.TemporaryDirectory(dir=os.environ["ARCHEAXIS_RUN_ROOT"])

    def test_chart_categories_and_values_are_projected_with_an_anchor(self):
        with self._tmp() as box:
            path = _pptx_with_chart(Path(box) / "chart.pptx")
            out = self.worker.extract(str(path))
            params = out["loss_receipt"]["params"]
            self.assertEqual(params["charts"], 1)
            self.assertEqual(params["charts_with_cached_values"], 1)
            self.assertEqual(params["charts_without_cached_values"], 0)
            self.assertIn("cached", params["chart_data_source"])
            self.assertIn("Q1=41", out["text"])
            self.assertIn("Q2=92.5", out["text"])
            self.assertIn("Revenue", out["text"])
            chart_parts = [item for item in out["structure"] if item["kind"] == "slide_chart"]
            self.assertEqual(len(chart_parts), 1)
            anchor = chart_parts[0]
            slice_text = out["text"][anchor["char_start"]:anchor["char_end"]]
            self.assertTrue(slice_text.startswith("Chart:"), slice_text)
            self.assertIn("Revenue", slice_text)
            self.assertIn("Q1=41", slice_text)
            self.assertEqual(anchor["path"], ["slide-1", "slide_chart"])
            self.assertGreaterEqual(anchor["char_start"], 0)
            self.assertLessEqual(anchor["char_end"], len(out["text"]))

    def test_a_series_without_cached_values_is_named_not_invented(self):
        with self._tmp() as box:
            source = _pptx_with_chart(Path(box) / "chart.pptx")
            stripped = _strip_caches(source, Path(box) / "linked.pptx")
            out = self.worker.extract(str(stripped))
            params = out["loss_receipt"]["params"]
            self.assertEqual(params["charts"], 1)
            self.assertEqual(params["charts_without_cached_values"], 1)
            self.assertEqual(params["charts_with_cached_values"], 0)
            self.assertIn("no cached values", out["text"])
            self.assertNotIn("Q1=41", out["text"])
            self.assertIn("never recomputed", out["loss_receipt"]["loss_note"])

    def test_a_package_without_charts_reports_no_chart_reads(self):
        with self._tmp() as box:
            from pptx import Presentation

            path = Path(box) / "plain.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.add_textbox(0, 0, 400000, 200000).text_frame.text = "no chart here"
            presentation.save(str(path))
            out = self.worker.extract(str(path))
            params = out["loss_receipt"]["params"]
            self.assertEqual(params["charts"], 0)
            self.assertEqual(params["charts_with_cached_values"], 0)
            self.assertNotIn("slide_chart", [item["kind"] for item in out["structure"]])


if __name__ == "__main__":
    unittest.main()
